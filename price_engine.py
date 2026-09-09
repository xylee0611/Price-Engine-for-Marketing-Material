"""Price list loading and tier/currency computation.

Column lookup is by header text (row 1), not fixed column letters, so the
engine keeps working if a price list's column order shifts or gains extra
columns. Multiple price list files (xlsx and/or csv) can be loaded together
and are merged by SKU Code, so a tier's source column can live in whichever
file actually has it.
"""
import csv
import math
import re

import openpyxl

# Direct-lookup tiers: display name -> exact header text expected in row 1.
DIRECT_TIER_HEADERS = {
    "RRP USD": "RRP (USD)",
    "RLP USD": "RLP (USD)",
    "RRP MYR": "RRP (MYR)",
    "R1 MYR": "P1 (MYR)",
    "R2 MYR": "P2 (MYR)",
    "R3 MYR": "P3 (MYR)",
    "RRP SGD": "RRP (SGD)",
    "RLP SGD": "RLP (SGD)",
    "RRP RMB": "RRP (RMB)",
    "RLP RMB": "RLP (RMB)",
    # From the supplementary "SKU Retail Recommended Price" list — confirmed
    # by the user to be USD-denominated (header text itself doesn't say).
    "LBL USD": "Loosing Line",
    "HBL USD": "Manager Price",
}

# Calculated tiers: display name -> (base header, cumulative -5% steps).
# Confirmed formulas only — never add an entry here without an explicit rule
# from the user (see feedback-pricing-safety-rules memory: never invent a
# calculation).
CALCULATED_TIER_DEFS = {
    "R1 USD": ("RLP (USD)", 1),
    "R2 USD": ("RLP (USD)", 2),
    "R3 USD": ("RLP (USD)", 3),
    "R4 USD": ("RLP (USD)", 4),
    "R5 USD": ("RLP (USD)", 5),
    "R6 USD": ("RLP (USD)", 6),
    "R4 MYR": ("RRP (MYR)", 4),
    "R5 MYR": ("RRP (MYR)", 5),
    "R6 MYR": ("RRP (MYR)", 6),
}

# Headers a "Custom tier" can be built on top of — any documented price column.
KNOWN_PRICE_HEADERS = sorted(
    set(DIRECT_TIER_HEADERS.values()) | {base for base, _ in CALCULATED_TIER_DEFS.values()}
)

SKU_CODE_HEADER = "SKU Code"
COMBINE_SKU_HEADER_PREFIX = "Combine SKU Code"

SKU_TOKEN_RE = re.compile(r"^\d{4,8}$")
CURRENCY_IN_HEADER_RE = re.compile(r"\(([A-Z]{2,4})\)")

# Display ordering for the tier dropdown: grouped by currency, then by tier rank.
CURRENCY_ORDER = ["USD", "MYR", "SGD", "RMB"]
TIER_RANK = {
    "RRP": 0, "RLP": 1, "R1": 2, "R2": 3, "R3": 4, "R4": 5, "R5": 6, "R6": 7,
    "LBL": 8, "HBL": 9,
}


def _tier_sort_key(tier_name):
    prefix, currency = tier_name.split()[0], tier_name.split()[-1]
    currency_idx = CURRENCY_ORDER.index(currency) if currency in CURRENCY_ORDER else len(CURRENCY_ORDER)
    return (currency_idx, TIER_RANK.get(prefix, 99), tier_name)


def round_up_half(x):
    return math.ceil(round(x * 2, 6)) / 2


def round_nearest_half(x):
    return round(round(x * 2, 6)) / 2


def round_up_int(x):
    return float(math.ceil(round(x, 6)))


def round_price(x, mode):
    if mode == "up_half":
        return round_up_half(x)
    if mode == "nearest_half":
        return round_nearest_half(x)
    if mode == "up_int":
        return round_up_int(x)
    if mode == "nearest_int":
        return float(round(x))
    return x  # "none"


def _display_name(path):
    """Filename for a plain path string or a Streamlit UploadedFile-like object."""
    name = getattr(path, "name", None)
    return name if name else str(path)


class PriceList:
    """Loads one or more price list files (xlsx and/or csv) and merges them
    by SKU Code, so a tier's source column can come from whichever file
    actually has it — e.g. a supplementary CSV adding LBL/HBL columns for
    the same SKUs as the main workbook."""

    def __init__(self, paths):
        if isinstance(paths, (str, bytes)) or hasattr(paths, "read"):
            paths = [paths]

        self.by_sku = {}
        self.combine_to_sku = {}
        self.available_headers = set()
        self.sources = []

        for path in paths:
            headers, rows = self._load_file(path)
            self.available_headers |= set(headers)
            self.sources.append(_display_name(path))
            for sku, combine, fields in rows:
                if sku in self.by_sku:
                    self.by_sku[sku].update(fields)
                else:
                    self.by_sku[sku] = dict(fields)
                if combine:
                    self.combine_to_sku.setdefault(combine, sku)

        self.row_count = len(self.by_sku)

    def _load_file(self, path):
        name = _display_name(path).lower()
        if name.endswith(".csv"):
            return self._load_csv(path)
        return self._load_xlsx(path)

    def _load_xlsx(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]

        headers = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col).value
            if val:
                headers[str(val).strip()] = col

        sku_col = headers.get(SKU_CODE_HEADER)
        combine_col = next(
            (c for h, c in headers.items() if h.startswith(COMBINE_SKU_HEADER_PREFIX)),
            None,
        )
        if sku_col is None:
            raise ValueError(f'Could not find a "{SKU_CODE_HEADER}" column in row 1 of {_display_name(path)}')

        rows = []
        for row in range(2, ws.max_row + 1):
            sku_val = ws.cell(row=row, column=sku_col).value
            sku = str(sku_val).strip() if sku_val is not None else ""
            if not sku:
                continue
            fields = {h: ws.cell(row=row, column=c).value for h, c in headers.items()}
            combine = None
            if combine_col is not None:
                combine_val = ws.cell(row=row, column=combine_col).value
                combine = str(combine_val).strip() if combine_val is not None else ""
            rows.append((sku, combine, fields))
        return list(headers.keys()), rows

    def _load_csv(self, path):
        if hasattr(path, "read"):
            path.seek(0)
            text = path.read()
            if isinstance(text, bytes):
                text = text.decode("utf-8-sig")
            lines = text.splitlines()
        else:
            with open(path, encoding="utf-8-sig") as f:
                lines = f.read().splitlines()

        reader = csv.DictReader(lines)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        if SKU_CODE_HEADER not in headers:
            raise ValueError(f'Could not find a "{SKU_CODE_HEADER}" column in the header row of {_display_name(path)}')
        combine_header = next((h for h in headers if h.startswith(COMBINE_SKU_HEADER_PREFIX)), None)

        rows = []
        for raw_row in reader:
            fields = {h.strip(): v for h, v in raw_row.items() if h}
            sku = str(fields.get(SKU_CODE_HEADER, "")).strip()
            if not sku:
                continue
            combine = str(fields.get(combine_header, "")).strip() if combine_header else None
            rows.append((sku, combine, fields))
        return headers, rows

    def lookup(self, sku_text):
        """Return the price-list record for a SKU, checking SKU Code then
        Combine SKU Code. None if not found in either."""
        record = self.by_sku.get(sku_text)
        if record is not None:
            return record, "sku_code"
        sku = self.combine_to_sku.get(sku_text)
        if sku is not None:
            return self.by_sku.get(sku), "combine_code"
        return None, None

    def available_tiers(self):
        """Tier display names that are actually usable against this price list,
        sorted by currency then tier rank (RRP, RLP, R1..R6)."""
        tiers = [name for name, header in DIRECT_TIER_HEADERS.items() if header in self.available_headers]
        tiers += [
            name for name, (base, _n) in CALCULATED_TIER_DEFS.items() if base in self.available_headers
        ]
        return sorted(tiers, key=_tier_sort_key)

    def available_base_headers(self):
        """Known price columns present in this price list, for the Custom tier base picker."""
        return [h for h in KNOWN_PRICE_HEADERS if h in self.available_headers]


def currency_for_tier(tier_name):
    if tier_name in DIRECT_TIER_HEADERS or tier_name in CALCULATED_TIER_DEFS:
        return tier_name.split()[-1]
    return None


def currency_from_header(header):
    m = CURRENCY_IN_HEADER_RE.search(header)
    return m.group(1) if m else "?"


def compute_raw_price(record, tier_name):
    """Return the raw (unrounded) numeric price for a SKU's price-list record
    and a chosen tier, or None if the record lacks the needed column."""
    if tier_name in DIRECT_TIER_HEADERS:
        header = DIRECT_TIER_HEADERS[tier_name]
        val = record.get(header)
        return float(val) if val is not None else None

    if tier_name in CALCULATED_TIER_DEFS:
        base_header, n = CALCULATED_TIER_DEFS[tier_name]
        base = record.get(base_header)
        if base is None:
            return None
        return float(base) * (0.95 ** n)

    return None  # unsupported tier/currency — caller must not invent a formula


def compute_price(record, tier_name, rounding_mode="up_half"):
    raw = compute_raw_price(record, tier_name)
    if raw is None:
        return None, None
    final = round_price(raw, rounding_mode)
    return raw, final


def discount_factor(pct_per_step, steps):
    """Cascading markdown, e.g. -5% x3 -> 0.95**3."""
    return (1 - pct_per_step / 100.0) ** steps


def compute_custom_raw_price(record, base_header, factor):
    """User-defined tier: base_header value x an explicit factor (either a
    cascading-discount factor from discount_factor(), or a direct markup
    multiplier like 1.2 for RRP x 1.2). This is never applied automatically —
    the user must explicitly choose the base column and factor each time;
    nothing here is guessed from patterns in the data."""
    base = record.get(base_header)
    if base is None:
        return None
    return float(base) * factor


def compute_custom_price(record, base_header, factor, rounding_mode="up_half"):
    raw = compute_custom_raw_price(record, base_header, factor)
    if raw is None:
        return None, None
    return raw, round_price(raw, rounding_mode)


def format_label(currency, value):
    return f"{currency} {value:.1f}"
