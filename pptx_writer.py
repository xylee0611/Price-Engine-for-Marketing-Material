from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from placement import find_clear_position, pptx_label_size_emu

GAP_EMU = 100000
DEFAULT_FONT_SIZE = Pt(18)


def overlay_pptx(original_pptx_path, output_pptx_path, occurrences, prices, currency, obstacles_by_slide, slide_w, slide_h):
    prs = Presentation(original_pptx_path)
    placed = []
    skipped = []

    for occ in occurrences:
        sku = occ["sku"]
        if sku not in prices:
            continue
        pptx_info = occ.get("pptx")
        if pptx_info is None:
            skipped.append({"sku": sku, "reason": "No position available in the PPTX."})
            continue

        slide_idx = pptx_info["slide_idx"]
        anchor = (
            pptx_info["left"],
            pptx_info["top"],
            pptx_info["left"] + pptx_info["width"],
            pptx_info["top"] + pptx_info["height"],
        )
        font_size_emu = pptx_info["font_size"] or DEFAULT_FONT_SIZE
        label = f"{currency} {prices[sku]:.1f}"

        label_w, label_h = pptx_label_size_emu(label, font_size_emu)
        obstacles = [
            (l, t, l + w, t + h) for (l, t, w, h) in obstacles_by_slide.get(slide_idx, [])
        ]

        rect = find_clear_position(anchor, label_w, label_h, obstacles, (slide_w, slide_h), GAP_EMU)
        if rect is None:
            skipped.append({"sku": sku, "reason": f"No clear space found near SKU on slide {slide_idx + 1}."})
            continue

        slide = prs.slides[slide_idx]
        tb = slide.shapes.add_textbox(
            Emu(int(rect[0])), Emu(int(rect[1])), Emu(int(rect[2] - rect[0])), Emu(int(rect[3] - rect[1]))
        )
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        run = tf.paragraphs[0].add_run()
        run.text = label
        run.font.size = font_size_emu
        run.font.name = pptx_info["font_name"] or "Arial"
        run.font.bold = False
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT

        placed.append({"sku": sku, "slide": slide_idx + 1, "label": label})

    prs.save(output_pptx_path)
    return placed, skipped
